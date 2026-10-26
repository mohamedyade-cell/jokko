"""Générateur du jeu de démonstration « FICTIF ».

Toutes les entreprises produites sont inventées : raison sociale préfixée
« [FICTIF] », NINEA commençant par 000, RCCM au code « FIC », e-mails en
@fictif.example. Le classeur généré déclare « Jeu de données = FICTIF » dans
l'onglet Paramètres : l'import marque alors chaque entreprise comme fictive et
refuse de mélanger ce jeu avec des données réelles.

Le générateur reproduit la structure attendue du classeur v2.4 et de radar.db
(voir docs/IMPORT.md), avec des cas particuliers utiles aux tests :
doublons NINEA, homonymes, alias moteur, RCCM/NAEMA inversés, saisies sur un
ancien ID.
"""
from __future__ import annotations

import random
import sqlite3
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

SEGMENTS = [
    # nom, import, export, intensité, fréquence, couverture, potentiel, vague, accroche
    ("Négoce import", 20, 0, 12, 12, 8, 8, 1, "Sécurisez vos règlements fournisseurs en devises et lissez votre coût d'achat."),
    ("Industrie exportatrice", 10, 20, 12, 10, 8, 8, 1, "Optimisez le rapatriement de vos recettes export."),
    ("BTP et infrastructures", 15, 0, 10, 6, 6, 6, 2, "Couvrez vos achats d'équipements libellés en devises."),
    ("Agro-industrie", 12, 12, 9, 9, 6, 6, 2, "Protégez vos marges des variations EUR/USD."),
    ("Mines et énergie", 15, 15, 15, 9, 10, 10, 1, "Gérez vos flux USD importants avec un desk dédié."),
    ("Services", 6, 4, 4, 6, 3, 4, 3, "Simplifiez vos transferts internationaux."),
    ("Banques et établissements financiers", 10, 10, 12, 12, 6, 6, 4, "Partenariat interbancaire."),
    ("Autres secteurs", 4, 2, 3, 3, 2, 3, 3, "Découvrez nos solutions de change."),
]
TAILLES = [("TPE", 1, 3, 3, 2), ("PME", 2, 8, 9, 5), ("Grande", 3, 15, 15, 10)]
REGIONS = ["Dakar", "Thiès", "Saint-Louis", "Kaolack", "Ziguinchor", "Diourbel"]
FORMES = ["SA", "SARL", "SUARL", "SAS"]
SYLLABES = ["Al", "Bé", "Ca", "Dio", "Fa", "Ga", "Ka", "Lo", "Ma", "Né", "Ou", "Sa", "Té", "Ya", "Zi"]
COMMERCIAUX = ["Commercial Fictif 1", "Commercial Fictif 2", "Commercial Fictif 3"]


def _nom(rng: random.Random, i: int) -> str:
    return f"[FICTIF] {''.join(rng.choice(SYLLABES) for _ in range(3))} {rng.choice(['Négoce', 'Industries', 'Services', 'Trading', 'Distribution'])} {i:04d}"


def generer(dossier: str | Path, n: int = 200, version: int = 1, seed: int = 42) -> tuple[Path, Path]:
    """Produit SUNU_FX_SALES_RADAR_FICTIF.xlsx et radar_fictif.db dans `dossier`.

    Version 2 : 10 nouvelles entreprises, ENT-00005 fusionnée dans ENT-00004
    par le moteur (alias), ENT-00010 retirée de l'onglet Base.
    """
    rng = random.Random(seed)
    dossier = Path(dossier)
    dossier.mkdir(parents=True, exist_ok=True)
    total = n + (10 if version >= 2 else 0)
    entreprises = []
    for i in range(1, total + 1):
        seg = SEGMENTS[(i * 7) % len(SEGMENTS)]
        taille = TAILLES[(i * 3) % len(TAILLES)][0]
        e = {
            "id": f"ENT-{i:05d}", "nom": _nom(rng, i), "sigle": f"FIC{i}",
            "ninea": f"000{i:06d}", "rccm": f"SN-FIC-20{10 + i % 15}-B-{1000 + i}",
            "naema": f"{46 + i % 5}.{10 + i % 80:02d}", "activite": "Activité fictive de démonstration",
            "forme": FORMES[i % len(FORMES)], "adresse": f"{i} rue Fictive", "ville": REGIONS[i % len(REGIONS)],
            "region": REGIONS[i % len(REGIONS)],
            "tel": f"+221 30 {100 + i // 100:03d} {i % 100:02d} {i % 97:02d}" if i % 9 else None,
            "email": f"contact{i}@fictif.example" if i % 4 else None,
            "contact": f"Contact Fictif {i}" if i % 3 else None, "fonction": "Directeur financier" if i % 3 else None,
            "segment": seg[0], "taille": taille, "effectif": 5 + (i * 37) % 900,
            "import": "Oui" if i % 2 else ("Non" if i % 5 else None), "export": "Oui" if i % 3 == 0 else "Non",
            "commercial": COMMERCIAUX[i % 3] if i % 4 == 0 else None,
            "statut": "Validée" if i % 10 else "À vérifier", "fiches": 1 + i % 3,
        }
        entreprises.append(e)
    # Cas particuliers
    entreprises[6]["ninea"] = entreprises[5]["ninea"]  # ENT-00007 et ENT-00006 : même NINEA → doublon fort
    entreprises[8]["nom"] = entreprises[7]["nom"]  # ENT-00009 / ENT-00008 : homonymes, même ville ?
    entreprises[8]["ville"] = entreprises[7]["ville"]
    entreprises[8]["ninea"] = "000999999"
    entreprises[11]["rccm"] = "46.69"  # RCCM contenant un code NAEMA → anomalie
    base = [e for e in entreprises if not (version >= 2 and e["id"] in {"ENT-00005", "ENT-00010"})]

    wb = Workbook()
    ws = wb.active
    ws.title = "Base"
    ws.append(["SUNU FX SALES RADAR — JEU DE DONNÉES FICTIF (démonstration)"])
    entete = ["ID", "Raison sociale", "Sigle", "NINEA", "RCCM", "Code NAEMA", "Activité", "Forme juridique", "Adresse",
              "Ville", "Région", "Téléphone", "E-mail", "Dirigeant", "Fonction", "Segment", "Taille", "Effectif",
              "Import", "Export", "Commercial", "Statut validation", "Nb fiches"]
    ws.append(entete)
    for e in base:
        ws.append([e["id"], e["nom"], e["sigle"], e["ninea"], e["rccm"], e["naema"], e["activite"], e["forme"],
                   e["adresse"], e["ville"], e["region"], e["tel"], e["email"], e["contact"], e["fonction"], e["segment"],
                   e["taille"], e["effectif"], e["import"], e["export"], e["commercial"], e["statut"], e["fiches"]])

    ct = wb.create_sheet("Call Tracker")
    ct.append(["ID entreprise", "Date", "Type", "Tentative", "Joint", "Décideur", "RDV", "Résultat",
               "Prochaine action", "Date prochaine action", "Commercial", "Commentaire"])
    d0 = date(2026, 9, 1)
    for k in range(1, 41):
        eid = f"ENT-{(k * 3) % n + 1:05d}"
        joint = k % 2 == 0
        ct.append([eid, d0 + timedelta(days=k), "Appel" if k % 5 else "E-mail", 1 + k % 3, "Oui" if joint else "Non",
                   "Oui" if joint and k % 4 == 0 else "Non", "Oui" if k % 8 == 0 else "Non",
                   "Intéressé" if joint else "Pas de réponse", "Rappeler", d0 + timedelta(days=k + 7),
                   COMMERCIAUX[k % 3], f"Commentaire fictif {k}"])
    # Saisie rattachée à l'ancien ID ENT-00005 (sera réalignée sur ENT-00004 en version 2)
    ct.append(["ENT-00005", date(2026, 9, 20), "Appel", 1, "Oui", "Non", "Non", "Rappeler", "Rappeler",
               date(2026, 10, 1), COMMERCIAUX[0], "Saisie sur ancien ID"])

    qf = wb.create_sheet("Qualification FX")
    qf.append(["ID entreprise", "Devises", "Volume annuel (FCFA)", "Fréquence", "Banque actuelle", "Besoin confirmé",
               "Spot", "Forward", "Flexiterme", "Transferts", "Import confirmé", "Export confirmé", "Potentiel", "Date qualification"])
    for k, eid in enumerate(["ENT-00004", "ENT-00008", "ENT-00012", "ENT-00016", "ENT-00020"]):
        qf.append([eid, "EUR, USD", 1_500_000_000 * (k + 1), ["Mensuelle", "Hebdomadaire", "Trimestrielle"][k % 3],
                   "Banque Fictive", "Oui" if k % 2 == 0 else "Non", "Oui", "Oui" if k % 2 == 0 else "Non", "Non", "Oui",
                   "Oui", "Non" if k % 2 else "Oui", ["Élevé", "Moyen", "Faible"][k % 3], date(2026, 9, 15 + k)])

    pp = wb.create_sheet("Pipeline")
    pp.append(["ID entreprise", "Opportunité", "Produit", "Étape", "Montant", "Devise", "Marge (pb)", "Probabilité",
               "Date ouverture", "Date clôture", "Commercial"])
    pp.append(["ENT-00004", "Couverture achats EUR", "Forward", "Proposition / cotation", 2_000_000, "EUR", 30, 0.6,
               date(2026, 9, 20), date(2026, 11, 30), COMMERCIAUX[1]])
    pp.append(["ENT-00008", "Transferts USD", "Transferts", "12. Gagnée", 500_000, "USD", 20, 1, date(2026, 9, 10),
               date(2026, 10, 10), COMMERCIAUX[2]])
    pp.append(["ENT-00012", "Spot EUR", "Spot", "Négociation", 300_000_000, "XOF", 25, 0.75, date(2026, 10, 1),
               date(2026, 12, 15), COMMERCIAUX[0]])

    co = wb.create_sheet("Corrections")
    co.append(["ID entreprise", "Champ", "Avant", "Après", "Auteur", "Date", "Motif"])
    co.append(["ENT-00003", "Téléphone", entreprises[2]["tel"], "+221 30 999 99 03", "Commercial Fictif 1",
               date(2026, 9, 5), "Numéro vérifié par appel"])
    co.append(["ENT-00012", "RCCM", "46.69", "SN-FIC-2019-B-1012", "Commercial Fictif 1", date(2026, 9, 6),
               "Code NAEMA saisi par erreur dans le RCCM"])

    dbl = wb.create_sheet("Doublons")
    dbl.append(["ID 1", "ID 2", "Critère", "Décision", "Commentaire"])
    dbl.append(["ENT-00006", "ENT-00007", "NINEA identique", "", ""])

    pa = wb.create_sheet("Paramètres")
    pa.append(["Paramètre", "Valeur"])
    for row in (["Jeu de données", "FICTIF"], ["Seuil A", 75], ["Seuil B", 55], ["Seuil C", 35],
                ["Objectif couverture", 0.5], ["Date de lancement", date(2026, 9, 1)], ["Vagues lancées", "1, 2"],
                ["Marge par défaut (pb)", 25]):
        pa.append(row)

    sc = wb.create_sheet("Scoring")
    sc.append(["Segment", "Import", "Export", "Intensité", "Fréquence", "Couverture", "Potentiel", "Vague", "Accroche"])
    for s in SEGMENTS:
        sc.append(list(s))
    sc.add_table(Table(displayName="Segments", ref=f"A1:I{len(SEGMENTS) + 1}"))
    start = len(SEGMENTS) + 4
    sc.cell(row=start, column=1, value="Taille")
    for j, h in enumerate(["Ordre", "Intensité", "Fréquence", "Potentiel"], start=2):
        sc.cell(row=start, column=j, value=h)
    for k, t in enumerate(TAILLES, start=1):
        for j, v in enumerate(t, start=1):
            sc.cell(row=start + k, column=j, value=v)
    sc.add_table(Table(displayName="Tailles", ref=f"A{start}:E{start + len(TAILLES)}"))

    xlsx = dossier / "SUNU_FX_SALES_RADAR_FICTIF.xlsx"
    wb.save(xlsx)

    radar = dossier / "radar_fictif.db"
    if radar.exists():
        radar.unlink()
    con = sqlite3.connect(radar)
    con.executescript("""
        CREATE TABLE files (id INTEGER PRIMARY KEY, filename TEXT, pages INTEGER);
        CREATE TABLE sources (id INTEGER PRIMARY KEY, file_id INTEGER, page INTEGER, raison_sociale TEXT,
                              ninea TEXT, rccm TEXT, telephone TEXT, email TEXT, segment TEXT);
        CREATE TABLE entity_map (source_id INTEGER, entity_id TEXT);
        CREATE TABLE entity_alias (old_id TEXT, new_id TEXT, reason TEXT);
    """)
    nb_pdf = max(3, total // 40)
    con.executemany("INSERT INTO files VALUES (?,?,?)", [(f, f"CCI_FICTIF_{f:03d}.pdf", 40) for f in range(1, nb_pdf + 1)])
    sid = 0
    for e in entreprises:
        for k in range(e["fiches"]):
            sid += 1
            con.execute("INSERT INTO sources VALUES (?,?,?,?,?,?,?,?,?)",
                        (sid, 1 + sid % nb_pdf, 1 + sid % 40, e["nom"], e["ninea"], e["rccm"],
                         e["tel"] if k == 0 else None, e["email"], e["segment"]))
            target = "ENT-00004" if (version >= 2 and e["id"] == "ENT-00005") else e["id"]
            con.execute("INSERT INTO entity_map VALUES (?,?)", (sid, target))
    if version >= 2:
        con.execute("INSERT INTO entity_alias VALUES ('ENT-00005','ENT-00004','Fusion moteur (fictif)')")
    con.commit()
    con.close()
    return xlsx, radar
