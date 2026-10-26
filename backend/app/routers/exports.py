"""Exports Excel / CSV de toute vue filtrée et rapport hebdomadaire PDF."""
from __future__ import annotations

import csv
import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from .. import kpi
from ..db import get_db
from ..models import ROLE_FX_SALES, ROLE_RESPONSABLE, Entreprise, Utilisateur
from ..security import lecteur
from ..utils import date_fr
from . import activite, admin, entreprises, pilotage
from .commun import TRIS, filtrer_entreprises, item_entreprise

router = APIRouter(prefix="/api/exports", tags=["Exports"])

COLONNES = {
    "entreprises": [("id", "ID"), ("raison_sociale", "Raison sociale"), ("sigle", "Sigle"), ("ninea", "NINEA"),
                    ("rccm", "RCCM"), ("naema_code", "Code NAEMA"), ("telephone", "Téléphone"), ("ville", "Ville"),
                    ("region", "Région"), ("segment", "Segment"), ("taille", "Taille"), ("score", "Score"),
                    ("categorie", "Catégorie"), ("vague", "Vague"), ("statut_validation", "Validation"),
                    ("commercial", "Commercial"), ("fictif", "Fictif")],
    "interactions": [("date", "Date"), ("entreprise_id", "ID"), ("raison_sociale", "Raison sociale"), ("type", "Type"),
                     ("tentative", "Tentative"), ("joint", "Joint"), ("decideur_joint", "Décideur"), ("rdv_obtenu", "RDV"),
                     ("date_rdv", "Date RDV"), ("resultat", "Résultat"), ("prochaine_action", "Prochaine action"),
                     ("date_prochaine_action", "Date prochaine action"), ("commercial", "Commercial"), ("commentaire", "Commentaire")],
    "opportunites": [("entreprise_id", "ID"), ("raison_sociale", "Raison sociale"), ("intitule", "Opportunité"),
                     ("produit", "Produit"), ("etape", "Étape"), ("montant", "Montant"), ("devise", "Devise"),
                     ("montant_fcfa", "Montant FCFA"), ("marge_pb", "Marge (pb)"), ("pnb_estime_fcfa", "PNB estimé FCFA"),
                     ("probabilite", "Probabilité %"), ("date_ouverture", "Ouverture"), ("date_cloture_prevue", "Clôture prévue"),
                     ("commercial", "Commercial")],
    "agenda": [("rubrique", "Rubrique"), ("date_prochaine_action", "Date relance"), ("date_rdv", "Date RDV"),
               ("entreprise_id", "ID"), ("raison_sociale", "Raison sociale"), ("telephone", "Téléphone"),
               ("prochaine_action", "Action"), ("commercial", "Commercial")],
    "wallet": [("commercial", "Commercial"), ("portefeuille", "Portefeuille"), ("contactees", "Contactées"),
               ("jointes", "Jointes"), ("decideur", "Décideur joint"), ("qualifiees", "Qualifiées"), ("rdv", "RDV"),
               ("opportunites", "Opportunités"), ("clientes", "Clientes"), ("interactions", "Interactions"),
               ("montant_pipeline_fcfa", "Pipeline FCFA"), ("taux_couverture", "Couverture %")],
    "anomalies": [("entreprise_id", "ID"), ("raison_sociale", "Raison sociale"), ("type", "Anomalie"),
                  ("champ", "Champ"), ("valeur", "Valeur")],
    "corrections": [("date", "Date"), ("entreprise_id", "ID"), ("champ", "Champ"), ("avant", "Avant"),
                    ("apres", "Après"), ("auteur", "Auteur"), ("motif", "Motif"), ("origine", "Origine")],
    "audit": [("horodatage", "Horodatage"), ("utilisateur", "Utilisateur"), ("action", "Action"), ("table", "Table"),
              ("enregistrement_id", "Enregistrement"), ("ancienne_valeur", "Ancienne valeur"), ("nouvelle_valeur", "Nouvelle valeur")],
}


def _valeur(v):
    if isinstance(v, bool):
        return "Oui" if v else "Non"
    if isinstance(v, str) and len(v) >= 10 and v[4] == "-" and v[7] == "-":
        try:
            return date_fr(datetime.fromisoformat(v)) + (v[10:].replace("T", " ") if len(v) > 10 else "")
        except ValueError:
            return v
    if isinstance(v, (dict, list)):
        return str(v)
    if isinstance(v, str) and (v[:1] in "=@\t\r" or (v[:1] in "+-" and not v[1:2].isdigit())):
        return "'" + v  # neutralise l'injection de formules dans Excel
    return v


def _lignes(vue: str, request: Request, user: Utilisateur, db: Session) -> list[dict]:
    p = dict(request.query_params)
    p.pop("format", None)
    if vue == "entreprises":
        q = filtrer_entreprises(db, user, q=p.get("q"), segment=p.get("segment"), categorie=p.get("categorie"),
                                vague=int(p["vague"]) if p.get("vague") else None, taille=p.get("taille"),
                                region=p.get("region"), commercial_id=int(p["commercial_id"]) if p.get("commercial_id") else None,
                                statut=p.get("statut"), statut_validation=p.get("statut_validation"),
                                non_attribuees=p.get("non_attribuees") == "true")
        return [item_entreprise(e) for e in q.order_by(TRIS.get(p.get("tri", "score"), TRIS["score"]), Entreprise.id)]
    if vue == "interactions":
        return activite.lister_interactions(entreprise_id=p.get("entreprise_id"),
                                            commercial_id=int(p["commercial_id"]) if p.get("commercial_id") else None,
                                            du=date.fromisoformat(p["du"]) if p.get("du") else None,
                                            au=date.fromisoformat(p["au"]) if p.get("au") else None, limite=5000, user=user, db=db)
    if vue == "opportunites":
        return activite.lister_opportunites(commercial_id=int(p["commercial_id"]) if p.get("commercial_id") else None,
                                            etape=p.get("etape"), user=user, db=db)
    if vue == "agenda":
        a = activite.agenda(commercial_id=int(p["commercial_id"]) if p.get("commercial_id") else None,
                            jours=int(p.get("jours", 14)), aujourdhui=None, limite=100000, user=user, db=db)
        out = []
        for rub, lib in (("en_retard", "En retard"), ("aujourdhui", "Aujourd'hui"), ("a_venir", "À venir"), ("rdv", "RDV")):
            out += [{**x, "rubrique": lib} for x in a[rub]]
        return out
    if vue == "wallet":
        return pilotage.wallet(user=user, db=db)["commerciaux"]
    if vue == "anomalies":
        return pilotage.anomalies(type=p.get("type"), limite=100000, user=user, db=db)["items"]
    if vue == "corrections":
        return pilotage.corrections(entreprise_id=p.get("entreprise_id"), limite=100000, user=user, db=db)
    if vue == "audit":
        if user.role != ROLE_RESPONSABLE:
            raise HTTPException(403, "Journal d'audit réservé au Responsable")
        return admin.journal(table=p.get("table"), utilisateur=p.get("utilisateur"), enregistrement_id=p.get("enregistrement_id"),
                             action=p.get("action"), limite=5000, user=user, db=db)
    raise HTTPException(404, "Vue inconnue")


def _fictif(db: Session) -> bool:
    return db.query(Entreprise.id).filter(Entreprise.fictif.is_(True)).first() is not None


@router.get("/{vue}")
def exporter(vue: str, request: Request, format: str = "xlsx", user: Utilisateur = Depends(lecteur),
             db: Session = Depends(get_db)) -> Response:
    if vue not in COLONNES:
        raise HTTPException(404, "Vue inconnue")
    lignes = _lignes(vue, request, user, db)
    cols = COLONNES[vue]
    horodatage = datetime.now().strftime("%Y%m%d_%H%M")
    fictif = _fictif(db)
    nom = f"radar_{vue}_{horodatage}{'_FICTIF' if fictif else ''}"
    pilotage.journal_export(db, vue, len(lignes))
    if format == "csv":
        buf = io.StringIO()
        buf.write("﻿")  # BOM : ouverture correcte des accents dans Excel
        w = csv.writer(buf, delimiter=";")
        if fictif:
            w.writerow(["DONNÉES FICTIVES — jeu de démonstration"])
        w.writerow([lib for _, lib in cols])
        for l in lignes:
            w.writerow([_valeur(l.get(k)) for k, _ in cols])
        return Response(buf.getvalue().encode("utf-8"), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": f'attachment; filename="{nom}.csv"'})
    if format != "xlsx":
        raise HTTPException(422, "Format attendu : xlsx ou csv")
    wb = Workbook()
    ws = wb.active
    ws.title = vue[:30]
    if fictif:
        ws.append(["DONNÉES FICTIVES — jeu de démonstration"])
        ws["A1"].font = Font(bold=True, color="C00000")
    ws.append([lib for _, lib in cols])
    hdr = ws.max_row
    for c in ws[hdr]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="C8102E")
    for l in lignes:
        ws.append([_valeur(l.get(k)) for k, _ in cols])
    for k, (cle, _) in enumerate(cols, start=1):
        if "fcfa" in cle or cle in {"montant"}:
            for row in ws.iter_rows(min_row=hdr + 1, min_col=k, max_col=k):
                row[0].number_format = "# ##0"
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    ws.auto_filter.ref = f"A{hdr}:{ws.cell(row=hdr, column=len(cols)).column_letter}{ws.max_row}"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{nom}.xlsx"'})


# --- Rapport hebdomadaire PDF ----------------------------------------------------------------------

def _fmt(n) -> str:
    if n is None:
        return "—"
    if isinstance(n, float) and not n.is_integer():
        return f"{n:,.1f}".replace(",", " ").replace(".", ",")
    return f"{int(n):,}".replace(",", " ")


def rapport_hebdo_pdf(db: Session, commercial_id: int | None = None, today: date | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    data = kpi.calculer(db, commercial_id=commercial_id, today=today)
    rouge = colors.HexColor("#C8102E")
    styles = getSampleStyleSheet()
    styles["Title"].textColor = rouge
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.6 * cm, rightMargin=1.6 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                            title="SUNU FX SALES RADAR — Rapport hebdomadaire")
    story = [Paragraph("SUNU FX SALES RADAR — Rapport hebdomadaire", styles["Title"]),
             Paragraph(f"Édité le {date_fr(date.fromisoformat(data['date']))} — périmètre : "
                       f"{'commercial' if commercial_id else 'équipe Capital Market'}", styles["Normal"])]
    if _fictif(db):
        story.append(Paragraph("<font color='#C00000'><b>DONNÉES FICTIVES — jeu de démonstration, aucune entreprise réelle.</b></font>", styles["Normal"]))
    story.append(Spacer(1, 10))

    def tableau(titre: str, lignes: list[list], largeurs=None) -> None:
        story.append(Paragraph(titre, styles["Heading2"]))
        t = Table(lignes, colWidths=largeurs, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), rouge), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                               ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                               ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F7F7")])]))
        story.append(t)
        story.append(Spacer(1, 8))

    v, c, t = data["volumes"], data["couverture"], data["taux"]
    tableau("Volumes et couverture", [
        ["Indicateur", "Valeur"], ["Fiches", _fmt(v["fiches"])], ["Entreprises uniques", _fmt(v["entreprises_uniques"])],
        ["Base validée", _fmt(v["base_validee"])], ["Couverture réelle", f"{_fmt(c['reelle_pct'])} %"],
        ["Objectif à date", f"{_fmt(c['attendu_a_date_pct'])} %"],
        [f"Objectif {_fmt(c['duree_mois'])} mois", f"{_fmt(c['objectif_pct'])} % ({_fmt(c['objectif_entreprises'])} entreprises)"],
        ["Taux de joignabilité", f"{_fmt(t['joignabilite'])} %"], ["Taux de RDV", f"{_fmt(t['rdv'])} %"],
        ["Taux de conversion", f"{_fmt(t['conversion'])} %"]], [9 * cm, 6 * cm])
    tableau("Entonnoir", [["Étape", "Entreprises"]] + [[e["libelle"], _fmt(e["valeur"])] for e in data["entonnoir"]], [9 * cm, 6 * cm])
    tableau("Activité des 12 dernières semaines", [["Semaine", "Interactions", "Contactées", "Jointes", "RDV", "Opportunités"]] +
            [[s["periode"], _fmt(s["interactions"]), _fmt(s["contactees"]), _fmt(s["jointes"]), _fmt(s["rdv"]), _fmt(s["opportunites"])]
             for s in data["hebdomadaire"]])
    tableau("Performance par commercial", [["Commercial", "Portef.", "Contact.", "Jointes", "RDV", "Opp.", "Clientes", "Pipeline FCFA"]] +
            [[p["commercial"], _fmt(p["portefeuille"]), _fmt(p["contactees"]), _fmt(p["jointes"]), _fmt(p["rdv"]),
              _fmt(p["opportunites"]), _fmt(p["clientes"]), _fmt(p["montant_pipeline_fcfa"])] for p in data["performance"]])
    pp = data["pipeline"]
    tableau("Pipeline", [["Indicateur", "Valeur"], ["Opportunités ouvertes", _fmt(pp["opportunites_ouvertes"])],
                         ["Montant ouvert", f"{_fmt(pp['montant_ouvert_fcfa'])} FCFA"],
                         ["Montant pondéré", f"{_fmt(pp['montant_pondere_fcfa'])} FCFA"],
                         ["PNB estimé", f"{_fmt(pp['pnb_estime_fcfa'])} FCFA"],
                         ["Opportunités gagnées", _fmt(pp["gagnees"])]], [9 * cm, 6 * cm])
    doc.build(story)
    return buf.getvalue()


@router.get("/rapport/hebdomadaire.pdf")
def rapport_hebdo(commercial_id: int | None = None, user: Utilisateur = Depends(lecteur), db: Session = Depends(get_db)) -> Response:
    if user.role == ROLE_FX_SALES:
        commercial_id = user.id
    pdf = rapport_hebdo_pdf(db, commercial_id)
    pilotage.journal_export(db, "rapport_hebdomadaire", 1)
    return Response(pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="rapport_hebdo_{date.today():%Y%m%d}.pdf"'})


__all__ = ["router", "entreprises"]
