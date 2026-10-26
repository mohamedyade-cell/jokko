"""Lecture du classeur SUNU_FX_SALES_RADAR.xlsx (lecture seule, valeurs calculées)."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..utils import cle_colonne
from . import colonnes as C


@dataclass
class Table:
    nom: str
    lignes: list[dict[str, Any]] = field(default_factory=list)  # chaque ligne : champs canoniques + "_ligne"
    non_reconnues: list[str] = field(default_factory=list)
    trouvee: bool = False


def _vide(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _lire_plage(rows: list[tuple], mapping: dict[str, list[str]], nom: str, premiere_ligne: int = 1,
                max_entete: int = 15) -> Table:
    idx = C.index_alias(mapping)
    table = Table(nom=nom)
    best, best_row = 0, None
    for i, row in enumerate(rows[:max_entete]):
        n = sum(1 for v in row if not _vide(v) and cle_colonne(v) in idx)
        if n > best:
            best, best_row = n, i
    if best_row is None or best < 2:
        return table
    table.trouvee = True
    entete = rows[best_row]
    colonnes: dict[int, str] = {}
    for j, v in enumerate(entete):
        if _vide(v):
            continue
        champ = idx.get(cle_colonne(v))
        if champ and champ not in colonnes.values():
            colonnes[j] = champ
        else:
            table.non_reconnues.append(str(v))
    for i, row in enumerate(rows[best_row + 1:], start=best_row + 1):
        if all(_vide(v) for v in row):
            continue
        ligne = {champ: (row[j] if j < len(row) else None) for j, champ in colonnes.items()}
        if all(_vide(v) for v in ligne.values()):
            continue
        ligne["_ligne"] = premiere_ligne + i
        table.lignes.append(ligne)
    return table


class Classeur:
    def __init__(self, chemin: str | Path):
        self.chemin = Path(chemin)
        # data_only=True : valeurs calculées par Excel, jamais les formules
        self.wb = load_workbook(self.chemin, read_only=False, data_only=True)
        self._feuilles = {cle_colonne(n): n for n in self.wb.sheetnames}

    def feuille(self, cle: str):
        for alias in C.ONGLETS[cle]:
            if alias in self._feuilles:
                return self.wb[self._feuilles[alias]]
        return None

    def table(self, cle: str, mapping: dict[str, list[str]]) -> Table:
        ws = self.feuille(cle)
        if ws is None:
            return Table(nom=cle)
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        return _lire_plage(rows, mapping, cle)

    def tables_scoring(self) -> tuple[Table, Table]:
        """Tables « Segments » et « Tailles » de l'onglet Scoring.

        Utilise les tableaux Excel nommés s'ils existent, sinon recherche les
        blocs dont l'en-tête commence par « Segment » / « Taille ».
        """
        ws = self.feuille("scoring")
        if ws is None:
            return Table("segments"), Table("tailles")
        res = {}
        for nom_table, ref in getattr(ws, "tables", {}).items():
            k = cle_colonne(nom_table)
            rows = [tuple(c.value for c in r) for r in ws[ref]]
            first_row = ws[ref][0][0].row
            if "segment" in k:
                res["segments"] = _lire_plage(rows, C.SEGMENTS, "segments", first_row, 2)
            elif "taille" in k:
                res["tailles"] = _lire_plage(rows, C.TAILLES, "tailles", first_row, 2)
        if len(res) < 2:
            all_rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
            for cle_bloc, mapping, mot in (("segments", C.SEGMENTS, "segment"), ("tailles", C.TAILLES, "taille")):
                if cle_bloc in res:
                    continue
                for i, row in enumerate(all_rows):
                    starts = [j for j, v in enumerate(row) if not _vide(v) and cle_colonne(v) == mot]
                    if starts:
                        j0 = starts[0]
                        bloc = []
                        for r in all_rows[i:]:
                            seg = r[j0:j0 + 12]
                            if all(_vide(v) for v in seg[:1]) and bloc:
                                break
                            bloc.append(seg)
                        # coupe le bloc à la première colonne vide de l'en-tête
                        width = next((k for k, v in enumerate(bloc[0]) if _vide(v)), len(bloc[0]))
                        bloc = [b[:width] for b in bloc]
                        res[cle_bloc] = _lire_plage(bloc, mapping, cle_bloc, i + 1, 1)
                        break
        return res.get("segments", Table("segments")), res.get("tailles", Table("tailles"))

    def parametres(self) -> tuple[dict[str, Any], list[str]]:
        """Onglet Paramètres sous forme clé / valeur (colonnes A et B, ou Paramètre / Valeur)."""
        ws = self.feuille("parametres")
        if ws is None:
            return {}, []
        idx = C.index_alias(C.PARAMETRES)
        out, inconnus = {}, []
        for row in ws.iter_rows(values_only=True):
            cells = [v for v in row if not _vide(v)]
            if len(cells) < 2:
                continue
            k = cle_colonne(cells[0])
            if k in {"parametre", "parametres", "cle"}:
                continue
            if k in idx:
                out[idx[k]] = cells[1]
            else:
                inconnus.append(str(cells[0]))
        return out, inconnus
