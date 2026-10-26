"""Import : aucun ID perdu, idempotence, alias, protection des saisies utilisateur."""
from __future__ import annotations

import pytest

from app.corrections import corriger
from app.importer.run import ImportRefuse, importer
from app.models import (
    STATUT_FUSIONNEE,
    Correction,
    Entreprise,
    EntiteAlias,
    Fiche,
    Interaction,
    Opportunite,
    ProvenanceChamp,
    Qualification,
)


def _compte(db):
    return {m.__name__: db.query(m).count() for m in (Entreprise, Interaction, Opportunite, Qualification, Correction, Fiche)}


def test_premier_import(db, imported):
    r = imported["resume"]
    assert r["nouvelles_entreprises"] == 200
    assert r["fiches"] == db.query(Fiche).count() > 200
    assert db.query(Entreprise).count() == 200
    assert all(e.fictif for e in db.query(Entreprise))
    assert imported["saisies"]["Call Tracker"]["importees"] == 41
    # anomalie RCCM/NAEMA signalée mais valeur jamais déplacée vers l'autre champ
    assert any(a["type"] == "rccm_naema" for a in imported["anomalies"])
    e12 = db.get(Entreprise, "ENT-00012")
    assert e12.rccm == "SN-FIC-2019-B-1012"  # correction du classeur appliquée
    assert e12.naema_code != e12.rccm


def test_reimport_idempotent(db, imported, demo_files):
    avant = _compte(db)
    ids_avant = {e.id for e in db.query(Entreprise)}
    r = importer(db, "resp", *demo_files[1])
    assert _compte(db) == avant
    assert {e.id for e in db.query(Entreprise)} == ids_avant
    assert r["resume"]["nouvelles_entreprises"] == 0
    assert r["saisies"]["Call Tracker"] == {"deja_presentes": 41}
    assert r["saisies"]["Pipeline"] == {"deja_presentes": 3}


def test_nouvelle_version_sans_perte_id(db, imported, demo_files):
    ids_v1 = {e.id for e in db.query(Entreprise)}
    nb_inter_ent5 = db.query(Interaction).filter_by(entreprise_id="ENT-00005").count()
    assert nb_inter_ent5 >= 1
    r = importer(db, "resp", *demo_files[2])
    ids_v2 = {e.id for e in db.query(Entreprise)}
    assert ids_v1 <= ids_v2, "aucun ID existant ne doit disparaître"
    assert len(ids_v2 - ids_v1) == 10 == r["resume"]["nouvelles_entreprises"]
    # Alias moteur ENT-00005 → ENT-00004 : fusion, saisies réalignées, rien supprimé
    e5 = db.get(Entreprise, "ENT-00005")
    assert e5.statut == STATUT_FUSIONNEE and e5.fusionnee_vers == "ENT-00004"
    assert db.query(Interaction).filter_by(entreprise_id="ENT-00005").count() == 0
    assert db.get(EntiteAlias, "ENT-00005").nouvel_id == "ENT-00004"
    assert {"ancien_id": "ENT-00005", "nouvel_id": "ENT-00004"}.items() <= r["ids_fusionnes"][0].items()
    # Entreprise retirée du classeur : conservée et signalée
    e10 = db.get(Entreprise, "ENT-00010")
    assert e10 is not None and e10.absente_dernier_import
    assert "ENT-00010" in r["absentes"]
    # Les fiches de l'ancien ID sont rattachées au nouvel ID
    assert db.query(Fiche).filter_by(entreprise_id="ENT-00005").count() == 0
    # Relance de la v2 : toujours idempotent
    avant = _compte(db)
    importer(db, "resp", *demo_files[2])
    assert _compte(db) == avant


def test_saisie_utilisateur_jamais_ecrasee(db, imported, demo_files, users):
    e = db.get(Entreprise, "ENT-00020")
    corriger(db, e, "telephone", "+221 77 000 00 20", "fx1", "Vérifié")
    e.commercial_id = users["fx2"].id
    e.champs_verrouilles = [*e.champs_verrouilles, "commercial_id"]
    q = db.get(Qualification, "ENT-00004")
    q.banque_actuelle = "Banque modifiée dans l'application"
    db.commit()
    r = importer(db, "resp", *demo_files[1])
    db.expire_all()
    e = db.get(Entreprise, "ENT-00020")
    assert e.telephone == "+221 77 000 00 20"
    assert e.commercial_id == users["fx2"].id
    assert db.get(Qualification, "ENT-00004").banque_actuelle == "Banque modifiée dans l'application"
    assert r["resume"]["valeurs_protegees_non_ecrasees"] >= 1


def test_provenance_par_champ(db, imported):
    p = db.query(ProvenanceChamp).filter_by(entreprise_id="ENT-00001", champ="ninea").one()
    assert p.fichier_id is not None and p.page is not None and p.exacte


def test_simulation_n_ecrit_rien(db, users, demo_files):
    rapport = importer(db, "resp", *demo_files[1], simulation=True)
    assert rapport["resume"]["nouvelles_entreprises"] == 200
    assert db.query(Entreprise).count() == 0


def test_refus_melange_fictif_reel(db, imported, demo_files, tmp_path):
    from openpyxl import load_workbook

    xlsx, _ = demo_files[1]
    wb = load_workbook(xlsx)
    ws = wb["Paramètres"]
    for row in ws.iter_rows():
        if row[0].value == "Jeu de données":
            row[1].value = "RÉEL"
    reel = tmp_path / "reel.xlsx"
    wb.save(reel)
    with pytest.raises(ImportRefuse):
        importer(db, "resp", reel, None)
    assert db.query(Entreprise).filter_by(fictif=False).count() == 0


def test_alias_applique_aux_saisies_du_classeur(db, users, demo_files):
    # Import direct de la v2 : la ligne Call Tracker sur ENT-00005 est rattachée à ENT-00004
    r = importer(db, "resp", *demo_files[2])
    assert r["resume"]["saisies_realignees_par_alias"] >= 1
    assert db.query(Interaction).filter_by(commentaire="Saisie sur ancien ID").one().entreprise_id == "ENT-00004"
    assert db.get(Entreprise, "ENT-00005") is None  # jamais créée par l'application
